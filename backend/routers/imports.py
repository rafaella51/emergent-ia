"""Importar planilha (Excel/CSV) de leads e deixar a Sofia recomendar o serviço de cada um."""
import base64
import csv
import io
import json
import re
import unicodedata
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query

from lib import bot
from lib.db import db
from models.schemas import ImportUpload, new_id, now_utc
from routers.auth import require_session

router = APIRouter(prefix="/import", tags=["importar"], dependencies=[Depends(require_session)])

MAX_ROWS = 500
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
SERVICOS = {"site", "gmn", "ambos", "nenhum"}


def _norm(txt: str) -> str:
    txt = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", txt.strip().lower())


# Palavras que identificam cada coluna, na ordem de preferência.
FIELDS = {
    "business": ["empresa", "negocio", "estabelecimento", "razao social", "nome fantasia", "title", "titulo", "loja"],
    "name": ["contato", "responsavel", "proprietario", "dono", "nome"],
    "email": ["e-mail", "email"],
    "phone": ["whatsapp", "celular", "telefone", "fone", "phone", "tel"],
    "website": ["website", "site", "url", "pagina"],
    "niche": ["nicho", "segmento", "categoria", "category", "ramo", "tipo"],
    "rating": ["nota", "rating", "estrelas", "stars"],
    "reviews": ["avaliacoes", "reviews", "qtd avaliacoes", "numero de avaliacoes"],
    "city": ["cidade", "bairro", "endereco", "address", "localizacao"],
}


def _map_columns(headers: List[str]) -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    normed = {h: _norm(h) for h in headers}
    for field, keys in FIELDS.items():
        for key in keys:
            hit = next((h for h, n in normed.items() if key in n and h not in mapping.values()), None)
            if hit:
                mapping[field] = hit
                break
    return mapping


def _read_rows(filename: str, raw: bytes) -> List[Dict[str, str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        hi = next((i for i, r in enumerate(rows[:10]) if sum(1 for c in r if c not in (None, "")) >= 2), 0)
        headers = [str(c).strip() if c not in (None, "") else f"Coluna {j + 1}" for j, c in enumerate(rows[hi])]
        out = []
        for r in rows[hi + 1:]:
            if not any(c not in (None, "") for c in r):
                continue
            out.append({headers[j]: ("" if c is None else str(c).strip()) for j, c in enumerate(r) if j < len(headers)})
        return out
    if name.endswith((".csv", ".txt")):
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        return [
            {k.strip(): (v or "").strip() for k, v in row.items() if k}
            for row in csv.DictReader(io.StringIO(text), dialect=dialect)
            if any((v or "").strip() for v in row.values())
        ]
    if name.endswith(".xls"):
        raise HTTPException(status_code=422, detail="Arquivo .xls antigo: abra no Excel e salve como .xlsx")
    raise HTTPException(status_code=422, detail="Envie um arquivo .xlsx ou .csv")


def _clean(doc: dict) -> dict:
    doc = dict(doc)
    doc.pop("_id", None)
    for k in ("created_at", "analyzed_at", "contacted_at"):
        if doc.get(k) is not None:
            doc[k] = doc[k].isoformat() if hasattr(doc[k], "isoformat") else doc[k]
    return doc


@router.post("/upload")
async def upload(payload: ImportUpload):
    try:
        raw = base64.b64decode(payload.data_base64.split(",")[-1])
    except Exception:  # noqa: BLE001
        raise HTTPException(status_code=422, detail="arquivo inválido")
    rows = _read_rows(payload.filename, raw)
    if not rows:
        raise HTTPException(status_code=422, detail="Não encontrei linhas com dados na planilha")
    rows = rows[:MAX_ROWS]
    mapping = _map_columns(list(rows[0].keys()))
    batch = new_id()
    docs = []
    for row in rows:
        get = lambda f: row.get(mapping[f], "") if f in mapping else ""  # noqa: E731
        email = get("email")
        if not EMAIL_RE.search(email or ""):
            found = EMAIL_RE.search(" ".join(row.values()))
            email = found.group(0) if found else ""
        else:
            email = EMAIL_RE.search(email).group(0)
        business = get("business") or get("name") or next((v for v in row.values() if v), "")
        docs.append({
            "id": new_id(), "batch": batch, "file": payload.filename,
            "business": business[:120], "name": (get("name") if get("name") != business else "")[:80],
            "email": email.lower(), "phone": get("phone")[:40], "website": get("website")[:200],
            "niche": get("niche")[:80], "rating": get("rating")[:10], "reviews": get("reviews")[:10],
            "city": get("city")[:120],
            "raw": {k[:40]: v[:200] for k, v in list(row.items())[:25]},
            "status": "pendente", "service": None, "priority": None, "reason": None,
            "created_at": now_utc(),
        })
    await db.imported_leads.insert_many(docs)
    return {
        "batch": batch, "total": len(docs),
        "with_email": sum(1 for d in docs if d["email"]),
        "columns": mapping,
    }


ANALYSIS_SYSTEM = """Você é a Sofia, analista comercial de uma freelancer que vende dois serviços:
- "site": criação de site ou landing page profissional;
- "gmn": otimização do perfil no Google Meu Negócio (fotos, avaliações, informações, ranqueamento local).

Para cada empresa da lista, decida o serviço mais indicado:
- "site": não tem site, ou o site é só rede social/link genérico (instagram, facebook, linktree, wa.me);
- "gmn": já tem site próprio, mas o perfil no Google é fraco (poucas avaliações, nota baixa, sem informações);
- "ambos": não tem site E o perfil no Google é fraco ou não dá pra saber;
- "nenhum": já tem site próprio e perfil forte (muitas avaliações e nota alta) — pouca oportunidade.
"prioridade" de 0 a 100: mais alta quando há uma falha clara (oportunidade) e um jeito de contato (e-mail/telefone).
"motivo": 1 frase curta e concreta em português, citando o dado que levou à decisão.
Responda SOMENTE com um JSON: [{"i": <número>, "servico": "...", "prioridade": <0-100>, "motivo": "..."}]"""


@router.post("/analyze")
async def analyze(limit: int = Query(default=15, ge=1, le=25)):
    pend = await db.imported_leads.find({"status": "pendente"}).sort("created_at", 1).to_list(limit)
    if not pend:
        left = 0
        return {"analyzed": 0, "remaining": left}
    items = []
    for i, d in enumerate(pend):
        info = {k: d.get(k) for k in ("business", "niche", "website", "rating", "reviews", "city") if d.get(k)}
        info["tem_email"] = bool(d.get("email"))
        info["tem_telefone"] = bool(d.get("phone"))
        extra = {k: v for k, v in (d.get("raw") or {}).items() if v and len(str(v)) < 120}
        items.append({"i": i, "dados": info, "planilha": extra})
    try:
        raw = await bot._gemini(ANALYSIS_SYSTEM, json.dumps(items, ensure_ascii=False), json_mode=True)
        txt = raw.strip()
        txt = txt[txt.find("["): txt.rfind("]") + 1]
        result = {int(r["i"]): r for r in json.loads(txt)}
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"A IA não conseguiu analisar agora: {str(exc)[:120]}")
    done = 0
    for i, d in enumerate(pend):
        r = result.get(i) or {}
        serv = str(r.get("servico", "")).lower().strip()
        serv = serv if serv in SERVICOS else "ambos"
        try:
            prio = max(0, min(100, int(r.get("prioridade", 50))))
        except (TypeError, ValueError):
            prio = 50
        await db.imported_leads.update_one({"id": d["id"]}, {"$set": {
            "status": "analisado", "service": serv, "priority": prio,
            "reason": str(r.get("motivo", ""))[:200], "analyzed_at": now_utc(),
        }})
        done += 1
    remaining = await db.imported_leads.count_documents({"status": "pendente"})
    return {"analyzed": done, "remaining": remaining}


@router.get("/leads")
async def list_leads(status: Optional[str] = Query(default=None)):
    q = {"status": status} if status else {}
    docs = await db.imported_leads.find(q).sort([("priority", -1), ("created_at", -1)]).to_list(1000)
    return [_clean(d) for d in docs]


@router.post("/clear")
async def clear(only_contacted: bool = Query(default=False)):
    q = {"status": "abordado"} if only_contacted else {}
    res = await db.imported_leads.delete_many(q)
    return {"deleted": res.deleted_count}
